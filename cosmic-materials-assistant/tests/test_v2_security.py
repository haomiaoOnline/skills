"""Security and third-party distribution regression using no project samples."""
import csv,hashlib,sys,tempfile,unittest
from datetime import date
from pathlib import Path
from openpyxl import load_workbook,Workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/"scripts"))
import v2_auth,v2_core,v2_io,v2_cli,v2_review

def sample_rows(root):
 source=root/"real_requirements.txt"
 source.write_text("订单创建与查询由客户输入请求后返回不同结果。\n",encoding="utf8")
 digest=v2_auth.sha_file(source)
 rows=[]
 for i in range(10):
  process="能力"+str(i//2)
  rows.append({
   "row_id":str(i+2),"level1":"订单业务","level2":"处理","level3":process,
   "functional_process":"处理"+process,"functional_user":"商户","trigger":"商户发起"+process,
   "subprocess":("接收" if i%2==0 else "输出")+str(i),"movement":"E" if i%2==0 else "X",
   "data_group":"对象"+str(i),"attributes":"字段"+str(i),"cfp":"1",
   "development_type":"新增" if i<2 else "复用" if i<9 else "利旧",
   "development_basis":"设计与复用证明"+str(i),"evidence_text":"订单创建与查询由客户输入请求后返回不同结果。",
   "source_locator":"real_requirements.txt#1",
   "source_file":"real_requirements.txt","source_sha256":digest,
   "review_status":"approved","reviewer":"审核负责人","reviewed_at":date.today().isoformat(),
   "boundary_basis":"边界设计文档第1节","movement_basis":"接口清单第"+str(i)+"项"
  })
 return rows

class ThirdPartySafety(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
  self.root=Path(self.tmp.name)

 def test_signature_binds_content_and_source_without_faking_human_identity(self):
  rows=sample_rows(self.root)
  keyfile=self.root/"reviewer-private.key"
  keyfile.write_bytes(b"reviewer-owned-key-material-must-be-kept-secret-"+b"A"*32)
  key=v2_auth.read_key(keyfile)
  unsigned=v2_core.validate(rows,approval_key=key,source_base=self.root)
  self.assertFalse(unsigned["passed"])
  self.assertEqual(unsigned["approved_cfp"],0)
  self.assertIn("UNVERIFIED_APPROVAL_SIGNATURE",{x["code"] for x in unsigned["issues"]})
  draft=self.root/"draft.xlsx"
  v2_io.export_workbook(rows,unsigned,draft)
  signed=self.root/"signed.xlsx"
  outcome=v2_review.sign_file(draft,signed,keyfile,source_root=self.root)
  self.assertEqual(outcome["approved_cfp"],10,outcome)
  complete=self.root/"complete.xlsx"
  result=v2_cli.run(signed,complete,approval_key_file=keyfile,source_root=self.root)
  self.assertEqual(result["state"],"READY_FOR_USER_ACCEPTANCE")
  self.assertEqual(result["approved_cfp"],10)
  wb=load_workbook(signed)
  page=wb["COSMIC工作量拆解"]
  headers={cell.value:cell.column for cell in page[1]}
  page.cell(2,headers["子过程"]).value="未重新批准的新子过程"
  tampered=self.root/"tampered.xlsx";wb.save(tampered)
  stale=v2_cli.run(tampered,self.root/"stale.xlsx",approval_key_file=keyfile,source_root=self.root)
  self.assertEqual(stale["state"],"BLOCKED_CANDIDATE_ONLY")
  self.assertEqual(stale["approved_cfp"],0)
  (self.root/"real_requirements.txt").write_text("被修改的需求\n",encoding="utf8")
  changed=v2_cli.run(signed,self.root/"changed.xlsx",approval_key_file=keyfile,source_root=self.root)
  self.assertEqual(changed["state"],"BLOCKED_CANDIDATE_ONLY")
  self.assertEqual(changed["approved_cfp"],0)

 def test_nonexistent_evidence_and_approval_strings_never_pass(self):
  rows=sample_rows(self.root)
  for r in rows:
   r["source_file"]="nonexistent.md"
   r["source_sha256"]="a"*64
   r["approval_signature"]="f"*64
  check=v2_core.validate(rows,source_base=self.root)
  self.assertFalse(check["passed"])
  self.assertEqual(check["approved_cfp"],0)
  self.assertIn("UNVERIFIABLE_SOURCE",{v["code"] for v in check["issues"]})

 def test_source_file_cannot_escape_authorized_source_root(self):
  allowed=self.root/"source_area";allowed.mkdir()
  rows=sample_rows(allowed)
  outside=self.root/"private-secret.txt";outside.write_text("Private",encoding="utf8")
  rows[0]["source_file"]=str(outside)
  rows[0]["source_sha256"]=v2_auth.sha_file(outside)
  self.assertFalse(v2_auth.source_valid(rows[0],allowed))
  self.assertFalse(v2_core.validate(rows,source_base=allowed)["passed"])

 def test_excel_formula_is_plain_text_and_workload_visible(self):
  rows=sample_rows(self.root)
  rows[0]["evidence_text"]='=HYPERLINK("https://invalid.example","OPEN")'
  doc=self.root/"formula.xlsx"
  v2_io.export_workbook(rows,v2_core.validate(rows),doc)
  page=load_workbook(doc)["COSMIC工作量拆解"]
  header={c.value:c.column for c in page[1]}
  attack=page.cell(2,header["需求原文依据"])
  self.assertEqual(attack.data_type,"s")
  self.assertTrue(attack.value.startswith("=HYPERLINK"))
  self.assertEqual(page.cell(2,header["折算系数"]).value,"1")
  self.assertEqual(page.cell(4,header["折算系数"]).value,"0.33")
  summary=load_workbook(doc)["待核查与规则说明"]
  self.assertIn("4.31",str([cell.value for row in summary for cell in row]))

 def test_row_drop_header_collision_ambiguous_sheet_and_pending(self):
  f=self.root/"partial.csv"
  f.write_text("一级模块,二级模块,三级模块,子过程,建设状态\n,,,,已建设\n订单,业务,新建,接收,E\n",encoding="utf8")
  records=v2_io.read_source(f)
  self.assertEqual(len(records),2)
  d=self.root/"duplicate.csv"
  d.write_text("一级模块,子过程,子过程\n服务,真动作,伪动作\n",encoding="utf8")
  with self.assertRaisesRegex(ValueError,"Duplicate"):
   v2_io.read_source(d)
  wb=Workbook();wb.active.title="旧案例";wb.active.append(["一级模块","二级模块","三级模块","子过程"])
  ws=wb.create_sheet("客户真实需求");ws.append(["一级模块","二级模块","三级模块","子过程"])
  multi=self.root/"multi.xlsx";wb.save(multi)
  with self.assertRaisesRegex(ValueError,"Ambiguous"):
   v2_io.read_source(multi)
  out=self.root/"contract.xlsx"
  pending=self.root/"contract.xlsx.pending";pending.write_text("Unrelated file",encoding="utf8")
  v2_io.export_workbook(sample_rows(self.root),v2_core.validate(sample_rows(self.root)),out)
  self.assertEqual(pending.read_text(encoding="utf8"),"Unrelated file")
  with self.assertRaises(FileExistsError):
   v2_io.export_workbook([],v2_core.validate([]),out)

if __name__=="__main__":
 unittest.main()
