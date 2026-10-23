import csv,sys,tempfile,unittest,json
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/"scripts"))
import v2_core,v2_io,v2_cli
class GenericAdapterTests(unittest.TestCase):
 def setUp(self):
  self.t=tempfile.TemporaryDirectory();self.addCleanup(self.t.cleanup)
  self.root=Path(self.t.name)
 def test_english_csv_alias_and_one_workbook(self):
  src=self.root/"ordinary_client.csv"
  with src.open("w",newline="",encoding="utf8") as f:
   w=csv.writer(f);w.writerow(["Level1","Level2","Level3","Process","Trigger","User","Subprocess","Movement","Data Group","Attributes","Development Type"])
   w.writerow(["Sales","Orders","Create","Create order","User submit","User","Receive order","E","Request","Order ID","New"])
  # Explicit map can attach arbitrary foreign column names.
  profile=self.root/"mapping.json"
  profile.write_text(json.dumps({"column_mapping":{"level1":"Level1","level2":"Level2","level3":"Level3","functional_process":"Process","subprocess":"Subprocess","movement":"Movement","data_group":"Data Group","attributes":"Attributes","development_type":"Development Type"}}))
  got=v2_io.read_source(src,column_mapping=json.loads(profile.read_text())["column_mapping"])
  self.assertEqual(got[0]["level1"],"Sales")
  out=self.root/"deliverable.xlsx"
  result=v2_cli.run(src,out,profile=profile)
  self.assertEqual(result["file_count"],1)
  self.assertTrue(out.is_file())
  self.assertEqual(sorted(p.name for p in self.root.glob("*.xlsx")),["deliverable.xlsx"])
  from openpyxl import load_workbook
  w=load_workbook(out)
  self.assertEqual(w.sheetnames,["COSMIC工作量拆解","待核查与规则说明"])
 def test_excel_header_order_and_names_independent(self):
  from openpyxl import Workbook
  w=Workbook();s=w.active;s.title="业务需求"
  s.append(["数据属性","数据组","三级模块","二级模块","一级模块","子过程描述","数据移动类型","功能过程"])
  s.append(["订单号","订单请求","订单管理","销售","商店","接收订单","E","提交订单"])
  f=self.root/"other.xlsx";w.save(f)
  records=v2_io.read_source(f)
  self.assertEqual(records[0]["subprocess"],"接收订单")
  self.assertEqual(records[0]["data_group"],"订单请求")
  self.assertEqual(records[0]["level3"],"订单管理")
 def test_header_can_start_after_cover_rows(self):
  from openpyxl import Workbook
  w=Workbook();sh=w.active;sh.title="需求表"
  sh.append(["业务功能清单 V2"]);sh.append(["请填入以下拆解"])
  sh.append(["功能过程","数据组","一级模块","二级模块","三级模块","子过程","数据移动类型"])
  sh.append(["创建订单","订单请求","项目","业务","订单","接收订单","E"])
  src=self.root/"cover.xlsx";w.save(src)
  parsed=v2_io.read_source(src)
  self.assertEqual(len(parsed),1)
  self.assertEqual(parsed[0]["row_id"],"4")
  self.assertEqual(parsed[0]["data_group"],"订单请求")
 def test_docx_text_has_no_fabricated_movements(self):
  from docx import Document
  f=self.root/"spec.docx"
  d=Document();d.add_paragraph("支持用户查询订单");d.save(f)
  rows=v2_io.read_source(f)
  self.assertEqual(rows[0].get("movement",""),"")
  self.assertNotIn("cfp",rows[0])
 def test_workbook_itself_can_be_reused_for_next_review_round(self):
  src=self.root/"start.csv"
  src.write_text("一级模块,二级模块,三级模块,功能过程,子过程描述,数据移动类型,数据组,数据属性\n系统,业务,模块,创建,接收请求,E,请求,编号\n",encoding="utf8")
  v1=self.root/"first.xlsx"
  result=v2_cli.run(src,v1)
  self.assertEqual(result["file_count"],1)
  v2=self.root/"second.xlsx"
  result2=v2_cli.run(v1,v2)
  self.assertEqual(result2["file_count"],1)
  self.assertTrue(v2.exists())
  self.assertEqual(v2_io.read_source(v2)[0]["subprocess"],"接收请求")
 def test_unknown_column_mapping_fails_closed(self):
  f=self.root/"q.csv";f.write_text("Unknown\nAnything\n",encoding="utf8")
  with self.assertRaisesRegex(ValueError,"unavailable"):
   v2_io.read_source(f,column_mapping={"data_group":"Missing"})
if __name__=="__main__":unittest.main()
