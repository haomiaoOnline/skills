#!/usr/bin/env python3
"""COSMIC v2 generic semantic file adapters and single-workbook export.

No Shandong project/column assumptions. No audit sidecars as deliverables.
The workbook contains the details plus status and evidence issues in tabs.
"""
from __future__ import annotations
import csv,json,re
import os,tempfile
from pathlib import Path
from typing import Mapping
from xml.etree import ElementTree
import v2_core

SEMANTIC_COLUMNS={
 "row_id":"来源行",
 "level1":"一级模块","level2":"二级模块","level3":"三级模块",
 "functional_user":"功能用户","functional_process":"功能过程",
 "functional_process_id":"功能过程ID","fur_id":"需求ID",
 "trigger":"触发事件","subprocess":"子过程","movement":"数据移动类型",
 "data_group":"数据组","attributes":"数据属性","cfp":"标准CFP",
 "development_type":"开发类型","development_basis":"开发类型依据",
 "evidence_text":"需求原文依据","source_locator":"来源定位","status":"建设状态",
 "review_status":"人工审核状态","reviewer":"审核人","reviewed_at":"审核日期",
 "boundary_basis":"软件边界依据","movement_basis":"数据移动依据",
 "source_file":"原始证据文件","source_sha256":"原始文件SHA256","approval_signature":"审批签名",
}
# Alternative vocabulary recognized without hardcoded Excel letters.
ALIASES={
 "row_id":("row_id","行号","编号","序号"),
 "level1":("level1","一级模块","一级功能模块","Level1","L1"),
 "level2":("level2","二级模块","二级功能模块","Level2","L2"),
 "level3":("level3","三级模块","三级功能模块","Level3","L3"),
 "functional_user":("functional_user","功能用户","Functional User","User"),
 "functional_process":("functional_process","功能过程","功能处理","Process","Functional Process"),
 "functional_process_id":("functional_process_id","功能过程ID","逻辑过程ID","Process ID"),
 "trigger":("trigger","触发事件","触发条件","Trigger","Trigger Event"),
 "subprocess":("subprocess","子过程","子过程描述","数据移动描述","Subprocess","Subprocess Description"),
 "movement":("movement","数据移动类型","数据移动","Movement","Movement Type"),
 "data_group":("data_group","数据组","数据对象","兴趣对象","Data Group","OOI"),
 "attributes":("attributes","数据属性","数据字段","Attributes","Data Attributes"),
 "cfp":("cfp","标准CFP","CFP","COSMIC CFP"),
 "development_type":("development_type","开发类型","功能点类型","Development Type"),
 "development_basis":("development_basis","开发类型依据","复用依据","Development Evidence"),
 "evidence_text":("evidence_text","需求原文依据","需求描述","功能用户需求","功能性用户需求","FUR","Requirement"),
 "source_locator":("source_locator","来源定位","来源","证据位置","Source Locator"),
 "status":("status","建设状态","建设情况","是否建设","Status"),
 "fur_id":("fur_id","需求ID","需求编号","FUR ID"),
 "review_status":("review_status","人工审核状态","复核状态","Review Status"),
 "reviewer":("reviewer","审核人","Reviewer"),
 "reviewed_at":("reviewed_at","审核日期","Review Date"),
 "boundary_basis":("boundary_basis","软件边界依据","边界依据","Boundary Basis"),
 "movement_basis":("movement_basis","数据移动依据","移动依据","Movement Basis"),
 "source_file":("source_file","原始证据文件","Source File"),
 "source_sha256":("source_sha256","原始文件SHA256","Source SHA256"),
 "approval_signature":("approval_signature","审批签名","Approval Signature"),
}
REQUIRED_IDENTIFICATION=("level1","level2","level3","subprocess")
COPY_DOWN=("level1","level2","level3","functional_process","trigger","functional_user")
def label(x):
 return re.sub(r"[^\w\u4e00-\u9fff]+","",str(x or "").casefold())

def mapping_for(headers,overrides=None):
 overrides=overrides or {}
 if not set(overrides)<=set(SEMANTIC_COLUMNS):
  raise ValueError("Unknown semantic mappings: "+repr(sorted(set(overrides)-set(SEMANTIC_COLUMNS))))
 lookup={label(h):str(h) for h in headers if h is not None}
 nonempty=[label(h) for h in headers if label(h)]
 if len(nonempty)!=len(set(nonempty)):
  raise ValueError("Duplicate normalized column headers; cannot infer trustworthy data")
 mapping={}
 for sem,alias in ALIASES.items():
  value=overrides.get(sem)
  if value:
   if str(value) in headers:mapping[sem]=str(value)
   elif label(value) in lookup:mapping[sem]=lookup[label(value)]
   else:raise ValueError("Configured column unavailable: "+sem+" = "+str(value))
  else:
   for a in alias:
    if label(a) in lookup:
     mapping[sem]=lookup[label(a)]
     break
 if not any(k in mapping for k in REQUIRED_IDENTIFICATION):
  raise ValueError("Input lacks COSMIC semantic columns; configure column_mapping")
 return mapping

def row_records(headers,body,source,mapping=None):
 mapping=mapping_for(headers,mapping)
 last={key:"" for key in COPY_DOWN}
 records=[]
 for line,values in body:
  if isinstance(values,dict):
   data={key:str(values.get(column,"") or "").strip() for key,column in mapping.items()}
  else:
   d={str(k):values[i] if i<len(values) else None for i,k in enumerate(headers)}
   data={key:str(d.get(column,"") or "").strip() for key,column in mapping.items()}
  if not any(data.values()):continue
  # Do not discard incomplete data: review must identify missing mandatory values.
  for key in COPY_DOWN:
   if data.get(key):last[key]=data[key]
   elif key in mapping:data[key]=last[key]
  data["row_id"]=data.get("row_id") or str(line)
  data["source_locator"]=data.get("source_locator") or str(source)+"#row="+str(line)
  data.setdefault("status","已建设")
  records.append(data)
 return records

def read_source(source,sheet=None,column_mapping=None):
 source=Path(source);suffix=source.suffix.lower()
 if suffix==".csv":
  with source.open(encoding="utf-8-sig",newline="") as fh:
   rows=list(csv.reader(fh))
  if not rows:raise ValueError("Empty CSV")
  return row_records(rows[0],((i,r) for i,r in enumerate(rows[1:],2)),
                     source,column_mapping)
 if suffix in (".xlsx",".xlsm"):
  try:from openpyxl import load_workbook
  except ImportError as exc:raise RuntimeError("Install openpyxl for XLSX") from exc
  wb=load_workbook(source,read_only=False,data_only=False)
  if sheet and sheet not in wb.sheetnames:raise ValueError("Sheet not found: "+sheet)
  candidates=[]
  for page in ([wb[sheet]] if sheet else wb):
   for row_number in range(1,min(page.max_row,20)+1):
    sample=[str(c.value or "") for c in page[row_number]]
    try:
     m=mapping_for(sample,column_mapping)
     if len(m)>=2:
      candidates.append((len(m),-row_number,page,row_number,sample))
    except ValueError:pass
  if not candidates:raise ValueError("No sheet/header with recognizable semantic columns")
  found_sheets={candidate[2].title for candidate in candidates}
  if len(found_sheets)>1:
   raise ValueError("Ambiguous candidate worksheets: "+repr(sorted(found_sheets))+"; specify --sheet")
  _,_,ws,header_row,header=max(candidates,key=lambda x:(x[0],x[1]))
  body=[(i,[c.value for c in row]) for i,row in enumerate(ws.iter_rows(min_row=header_row+1),header_row+1)]
  return row_records(header,body,source.name+"#sheet="+ws.title,column_mapping)
 if suffix in (".md",".txt",".docx"):
  if suffix==".docx":
   try:from docx import Document
   except ImportError as exc:raise RuntimeError("Install python-docx for DOCX") from exc
   doc=Document(source)
   lines=[p.text.strip() for p in doc.paragraphs if p.text.strip()]
   for table in doc.tables:
    lines.extend(" | ".join(c.text.strip() for c in r.cells) for r in table.rows)
  else:lines=source.read_text(encoding="utf-8-sig").splitlines()
  # Evidence intake only. No inferred E/R/W/X, CFP, or module identity.
  return [{"row_id":str(n),"source_locator":source.name+"#paragraph="+str(n),
           "evidence_text":text.strip(),"status":"待确认"}
          for n,text in enumerate(lines,1) if text.strip()]
 raise ValueError("Supported inputs: CSV, XLSX, DOCX, Markdown, TXT")

def export_workbook(records,validation,output):
 try:
  from openpyxl import Workbook
  from openpyxl.styles import Font,Alignment,PatternFill
  from openpyxl.utils import get_column_letter
 except ImportError as exc:raise RuntimeError("Install openpyxl to write XLSX") from exc
 output=Path(output)
 if output.suffix.lower()!=".xlsx":raise ValueError("Single deliverable must be an XLSX")
 if output.exists() or output.is_symlink():raise FileExistsError("Refusing to overwrite existing deliverable")
 wb=Workbook();sheet=wb.active;sheet.title="COSMIC工作量拆解"
 headers=list(SEMANTIC_COLUMNS)
 sheet.append([SEMANTIC_COLUMNS[h] for h in headers]+["折算系数","折算工作量"])
 for item in records:
  row=v2_core.record(item)
  weight=v2_core.WEIGHTS.get(row.get("development_type","")) if row.get("status") not in ("未建设","unbuilt","excluded") else None
  values=[row.get(k,"") for k in headers]
  values.extend([str(weight) if weight is not None else "",str(weight) if weight is not None and row.get("movement") in v2_core.MOVEMENTS else ""])
  sheet.append(values)
  for cell in sheet[sheet.max_row]:
   if cell.data_type=="f":cell.data_type="s" # never allow text input to become formula
 sheet.freeze_panes="A2"
 sheet.auto_filter.ref=sheet.dimensions
 for n in range(1,len(headers)+1):
  sheet.column_dimensions[get_column_letter(n)].width=23 if n<8 else 29
 sheet.column_dimensions[get_column_letter(headers.index("evidence_text")+1)].width=62
 sheet.column_dimensions[get_column_letter(headers.index("source_locator")+1)].width=45
 issues=wb.create_sheet("待核查与规则说明")
 issues.append(["级别","规则代码","来源行","字段","说明","关联行","规则性质"])
 for issue in validation["issues"]:
  issues.append(["阻断",issue["code"],issue["row_id"],issue["field"],
                 issue["message"],issue.get("related_row") or "",issue["basis"]])
 issues.append(["状态","交付准入",str(validation["passed"]),"","","",""])
 issues.append(["统计","候选数据移动",validation["candidate_cfp"],"","","",""])
 issues.append(["统计","可确认CFP",validation["approved_cfp"],"","","",""])
 issues.append(["统计","2:7:1目标",json.dumps(validation["ratio_target"],ensure_ascii=False),"","","",""])
 issues.append(["统计","2:7:1实际",json.dumps(validation["ratio_actual"],ensure_ascii=False),"","","",""])
 issues.append(["统计","复用系数",str(v2_core.WEIGHTS["复用"]),"","","",""])
 issues.append(["统计","组织折算工作量基数",validation["weighted_workload_basis"],"","","",""])
 issues.append(["说明","标准与客户规则",validation["disclaimer"],"","","",""])
 issues.column_dimensions["B"].width=36
 issues.column_dimensions["C"].width=40
 issues.column_dimensions["E"].width=86
 issues.freeze_panes="A2"
 for ws in (sheet,issues):
  for c in ws[1]:
   c.font=Font(bold=True,color="FFFFFF")
   c.fill=PatternFill("solid",fgColor="17324D")
  for row in ws.iter_rows(min_row=2):
   for c in row:c.alignment=Alignment(vertical="top",wrap_text=True)
 output.parent.mkdir(parents=True,exist_ok=True)
 fd,local_tmp=tempfile.mkstemp(prefix="cosmic-secure-",suffix=".xlsx",dir=output.parent)
 os.close(fd)
 temp=Path(local_tmp)
 try:
  wb.save(temp)
  os.link(temp,output) # refuse to overwrite an existing file, even if created concurrently
 finally:
  if temp.exists():temp.unlink()
 return output
